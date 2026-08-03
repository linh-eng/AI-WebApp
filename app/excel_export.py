"""Xuất báo cáo Excel: điền dữ liệu từ CSDL vào file mẫu Báo cáo Mua hàng THNG.

Nguyên tắc giữ nguyên định dạng: KHÔNG vẽ lại Excel. Ta mở chính file mẫu, chỉ
ghi giá trị vào các Ô NHẬP TAY theo đúng vị trí cột; mọi công thức (tên dự án,
giá trị, tiết kiệm, công nợ, tuổi nợ...) đã có sẵn trong file mẫu và tự tính lại
khi mở bằng Excel (đặt fullCalcOnLoad=True để buộc tính lại).
"""
from __future__ import annotations

import io

import openpyxl

from .config import TEMPLATE_XLSX

DATA_START_ROW = 4
MAX_ROWS = 200  # vùng công thức file mẫu: dòng 4 -> 203

# Ánh xạ (thuộc tính model -> chỉ số cột). CHỈ liệt kê ô nhập tay; bỏ qua cột
# công thức để không ghi đè công thức trong file mẫu.
COT_DU_AN = {
    "ma_du_an": 1, "ten": 2, "khach_hang": 3, "pm": 4, "ngay_bat_dau": 5,
    "ngay_ket_thuc": 6, "ngan_sach": 7, "trang_thai": 8, "ghi_chu": 9,
}
COT_NHA_CUNG_CAP = {
    "ma_ncc": 1, "ten": 2, "nguoi_lien_he": 3, "dien_thoai": 4, "email": 5,
    "dia_chi": 6, "dieu_khoan_tt": 7, "danh_gia": 8, "nhom_hang": 9,
    "ma_so_thue": 10, "ghi_chu": 11,
}
COT_CHECK_GIA = {
    "ma_yc": 1, "ngay_nhan": 2, "ma_du_an": 3, "nguoi_yc": 5, "hang_muc": 6,
    "dvt": 7, "so_luong": 8, "don_gia_du_toan": 9, "han_tra_gia": 11,
    "ngay_tra_gia": 12, "trang_thai": 14, "ghi_chu": 21,
}
COT_BAO_GIA = {
    "ma_bao_gia": 1, "ngay": 2, "ma_yc": 3, "ma_ncc": 9, "don_gia": 11,
    "thoi_gian_giao": 13, "hieu_luc": 15, "duoc_chon": 20, "ghi_chu": 22,
}
COT_PR = {
    "ma_pr": 1, "ngay": 2, "ma_yc": 3, "so_luong": 10, "don_gia_du_toan": 11,
    "ngay_can_hang": 13, "trang_thai_duyet": 14, "nguoi_duyet": 15,
    "ngay_duyet": 16, "ghi_chu": 17,
}
COT_PO = {
    "ma_po": 1, "ngay": 2, "ma_pr": 3, "ma_ncc": 8, "so_luong_dat": 10,
    "don_gia_po": 11, "so_hop_dong": 15, "ngay_ky": 16, "ngay_giao_cam_ket": 17,
    "ngay_nhan_thuc_te": 18, "sl_nhan": 21, "sl_loi": 22, "sl_tra_lai": 24,
    "ghi_chu": 28,
}
COT_THANH_TOAN = {
    "ma_phieu_chi": 1, "ngay": 2, "ma_po": 3, "so_tien": 8, "dot": 9,
    "hinh_thuc": 12, "so_chung_tu": 13, "ghi_chu": 14,
}


def _ghi_sheet(ws, records, colmap, max_rows=MAX_ROWS):
    """Xoá dữ liệu ví dụ ở các ô nhập tay rồi ghi records vào (giữ nguyên công thức)."""
    last_row = DATA_START_ROW + max_rows - 1
    for r in range(DATA_START_ROW, last_row + 1):
        for col in colmap.values():
            ws.cell(row=r, column=col).value = None
    for i, rec in enumerate(records[:max_rows]):
        r = DATA_START_ROW + i
        for attr, col in colmap.items():
            ws.cell(row=r, column=col).value = getattr(rec, attr)


def xuat_bao_cao(du_ans, nhas, check_gias, bao_gias, prs, pos, thanh_toans) -> bytes:
    """Trả về nội dung file .xlsx (bytes) đã điền dữ liệu, đúng mẫu THNG."""
    if not TEMPLATE_XLSX.exists():
        raise FileNotFoundError(f"Không tìm thấy file mẫu: {TEMPLATE_XLSX}")

    wb = openpyxl.load_workbook(TEMPLATE_XLSX)

    _ghi_sheet(wb["Dự án"], du_ans, COT_DU_AN)
    _ghi_sheet(wb["Nhà cung cấp"], nhas, COT_NHA_CUNG_CAP)
    _ghi_sheet(wb["Check giá"], check_gias, COT_CHECK_GIA)
    _ghi_sheet(wb["Báo giá"], bao_gias, COT_BAO_GIA)
    _ghi_sheet(wb["PR"], prs, COT_PR)
    _ghi_sheet(wb["PO"], pos, COT_PO)
    _ghi_sheet(wb["Thanh toán"], thanh_toans, COT_THANH_TOAN)

    # Buộc Excel tính lại toàn bộ công thức khi mở file.
    wb.calculation.fullCalcOnLoad = True

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()
